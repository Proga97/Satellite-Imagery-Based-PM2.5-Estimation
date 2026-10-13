#!/usr/bin/env python
"""Build docs/thesis/thesis.docx from the untouched Purdue template and the content
modules in docs/thesis/content/.

Every run starts from the template, so the output is deterministic and no edit
accumulates. Formatting follows the Purdue Standard Thesis Handbook (August 2026):
template styles only, captions as Word caption fields numbered chapter.sequence,
figure captions below, table captions above, alt text on every figure, real
sub/superscripts, bracketed numeric references with hyperlinks.

Text markup in content:  _{sub}  ^{sup}  *italic*  [@key] or [@a; @b] citations
                         {fig:key} {tab:key} {eq:key} cross-references
Once the docx is edited by hand in Word, stop running this script.
"""
import os
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "docs/thesis/content"))
from refs import REFS          # noqa: E402
import chapters as C           # noqa: E402

TEMPLATE = ROOT / "docs/thesis/Purdue-Standard-Thesis-Template-August-2026.docx"
OUT = (Path(os.environ["THESIS_OUT"]).resolve() if os.environ.get("THESIS_OUT") else ROOT / "docs/thesis/thesis.docx")
FIGDIR = ROOT / "docs/figures"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

CITE = re.compile(r"\[@([^\]]+)\]")
XREF = re.compile(r"\{(fig|tab|eq):([a-z0-9_]+)\}")
MARK = re.compile(r"_\{([^}]*)\}|\^\{([^}]*)\}|\*([^*]+)\*")
XNAME = {"fig": "Figure", "tab": "Table", "eq": "Equation"}


# ------------------------------------------------------------------ numbering
class Numbering:
    def __init__(self):
        self.cites, self.labels = {}, {}

    def scan_cites(self, text):
        for m in CITE.finditer(text or ""):
            for k in m.group(1).split(";"):
                k = k.strip().lstrip("@")
                if k not in REFS:
                    raise KeyError(f"unknown citation key: {k}")
                self.cites.setdefault(k, len(self.cites) + 1)

    def resolve(self, text):
        def cite(m):
            nums = sorted(self.cites[k.strip().lstrip("@")] for k in m.group(1).split(";"))
            return ", ".join(f"[{n}]" for n in nums)

        def xref(m):
            return f"{XNAME[m.group(1)]} {self.labels[(m.group(1), m.group(2))]}"
        return XREF.sub(xref, CITE.sub(cite, text))


NUM = Numbering()


def block_texts(b):
    k = b[0]
    if k in ("p", "h2", "h3"):
        return [b[1]]
    if k == "fig":
        return [b[3]]
    if k == "tab":
        d = b[2]
        return [d["caption"]] + list(d["header"]) + [c for r in d["rows"] for c in r] + [d.get("note") or ""]
    return []


def prepass():
    for i, (title, blocks) in enumerate(C.CHAPTERS, start=1):
        n = {"fig": 0, "tab": 0, "eq": 0}
        for b in blocks:
            if b[0] in n:
                n[b[0]] += 1
                key = b[1]
                if (b[0], key) in NUM.labels:
                    raise KeyError(f"duplicate label {b[0]}:{key}")
                NUM.labels[(b[0], key)] = f"{i}.{n[b[0]]}"
            for t in block_texts(b):
                NUM.scan_cites(t)


# ------------------------------------------------------------------ XML pieces
def mk_run(text, vert=None, italic=False, bold=False, size=None, underline=False):
    r = OxmlElement("w:r")
    props = []
    if bold:
        props += [OxmlElement("w:b"), OxmlElement("w:bCs")]
    if italic:
        props += [OxmlElement("w:i"), OxmlElement("w:iCs")]
    if size:
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag); e.set(qn("w:val"), str(size)); props.append(e)
    if underline:
        e = OxmlElement("w:u"); e.set(qn("w:val"), "single"); props.append(e)
    if vert:
        e = OxmlElement("w:vertAlign"); e.set(qn("w:val"), vert); props.append(e)
    if props:
        rpr = OxmlElement("w:rPr")
        for e in props:
            rpr.append(e)
        r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(XML_SPACE, "preserve")
    r.append(t)
    return r


def runs_for(marked, **kw):
    out, pos = [], 0
    for m in MARK.finditer(marked):
        if m.start() > pos:
            out.append(mk_run(marked[pos:m.start()], **kw))
        if m.group(1) is not None:
            out.append(mk_run(m.group(1), vert="subscript", **kw))
        elif m.group(2) is not None:
            out.append(mk_run(m.group(2), vert="superscript", **kw))
        else:
            out.append(mk_run(m.group(3), italic=True, **kw))
        pos = m.end()
    if pos < len(marked):
        out.append(mk_run(marked[pos:], **kw))
    return out


def plain(marked):
    return MARK.sub(lambda m: m.group(1) or m.group(2) or m.group(3), marked)


def para(style_id, marked="", jc=None, before=None, ind=None):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style_id); ppr.append(ps)
    if before is not None:
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), str(before)); ppr.append(sp)
    if ind:
        e = OxmlElement("w:ind")
        for k, v in ind.items():
            e.set(qn(f"w:{k}"), str(v))
        ppr.append(e)
    if jc:
        j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)
    p.append(ppr)
    for r in runs_for(marked):
        p.append(r)
    return p


def fld(instr, cached):
    f = OxmlElement("w:fldSimple"); f.set(qn("w:instr"), instr)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rpr.append(OxmlElement("w:noProof")); r.append(rpr)
    t = OxmlElement("w:t"); t.text = cached; r.append(t); f.append(r)
    return f


def caption(kind, label, marked):
    """Word caption exactly as 'Insert Caption' writes it with 'include chapter number'."""
    ch, n = label.split(".")
    full = f"{kind} {label}. {plain(marked)}"
    p = para("CaptionFigure" if kind == "Figure" else "CaptionTable",
             jc="center" if len(full) <= 90 else None)   # one line: centered; longer: justified
    if kind == "Table":                                   # a table caption stays on the page of its table
        p.find(qn("w:pPr")).insert(1, OxmlElement("w:keepNext"))   # schema order: pStyle, keepNext, ..., jc
    p.append(mk_run(f"{kind} "))
    p.append(fld(" STYLEREF 1 \\s ", ch))
    p.append(mk_run("."))
    p.append(fld(f" SEQ {kind} \\* ARABIC \\s 1 ", n))
    for r in runs_for(". " + marked):
        p.append(r)
    return p


_pic_id = [1000]


def figure_para(doc, fname, alt, max_w=6.2, max_h=6.9):
    path = FIGDIR / fname
    w, h = Image.open(path).size
    width = min(max_w, max_h * w / h)
    p_el = para("FigurePicture")
    p_el.find(qn("w:pPr")).insert(1, OxmlElement("w:keepNext"))   # the picture stays on the page of its caption
    Paragraph(p_el, doc._body).add_run().add_picture(str(path), width=Inches(width))
    _pic_id[0] += 1
    for el in p_el.iter():
        if el.tag.endswith("}docPr") or el.tag.endswith("}cNvPr"):
            el.set("id", str(_pic_id[0])); el.set("name", f"Picture {_pic_id[0]}"); el.set("descr", alt)
    return p_el


def table_el(header, rows, widths=None, size=20, align=None):
    ncol = len(header)
    widths = widths or [100 / ncol] * ncol
    assert abs(sum(widths) - 100) < 0.5 and len(widths) == ncol, (widths, header)
    align = align or (["left"] + ["center"] * (ncol - 1))
    tbl = OxmlElement("w:tbl")
    pr = OxmlElement("w:tblPr")
    st = OxmlElement("w:tblStyle"); st.set(qn("w:val"), "TableGrid"); pr.append(st)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), "5000"); tw.set(qn("w:type"), "pct"); pr.append(tw)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pr.append(jc)
    look = OxmlElement("w:tblLook")
    for k, v in dict(val="04A0", firstRow="1", lastRow="0", firstColumn="1", lastColumn="0", noHBand="0", noVBand="1").items():
        look.set(qn(f"w:{k}"), v)
    pr.append(look); tbl.append(pr)
    grid = OxmlElement("w:tblGrid")
    for wd in widths:
        g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(round(9360 * wd / 100))); grid.append(g)
    tbl.append(grid)
    for ri, row in enumerate([header] + rows):
        assert len(row) == ncol, row
        tr = OxmlElement("w:tr"); trpr = OxmlElement("w:trPr")
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))     # header row marked for screen readers
        j = OxmlElement("w:jc"); j.set(qn("w:val"), "center"); trpr.append(j)
        tr.append(trpr)
        for ci, cell in enumerate(row):
            tc = OxmlElement("w:tc"); tcpr = OxmlElement("w:tcPr")
            cw = OxmlElement("w:tcW"); cw.set(qn("w:w"), str(round(widths[ci] * 50))); cw.set(qn("w:type"), "pct"); tcpr.append(cw)
            va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcpr.append(va)
            tc.append(tcpr)
            p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
            if ri < len(rows):                      # all rows but the last: keep with next
                ppr.append(OxmlElement("w:keepNext"))
            sp = OxmlElement("w:spacing")
            for k, v in dict(before="60", after="60", line="240", lineRule="auto").items():
                sp.set(qn(f"w:{k}"), v)
            ppr.append(sp)
            a = OxmlElement("w:jc"); a.set(qn("w:val"), align[ci]); ppr.append(a)
            p.append(ppr)
            for r in runs_for(NUM.resolve(cell), size=size, bold=(ri == 0)):
                p.append(r)
            tc.append(p); tr.append(tc)
        tbl.append(tr)
    return tbl


def equation(parts, label, tabs):
    M = "Cambria Math"

    def mrun(text):
        r = OxmlElement("m:r"); rpr = OxmlElement("w:rPr"); f = OxmlElement("w:rFonts")
        f.set(qn("w:ascii"), M); f.set(qn("w:hAnsi"), M); rpr.append(f); r.append(rpr)
        t = OxmlElement("m:t"); t.text = text; t.set(XML_SPACE, "preserve"); r.append(t)
        return r

    def script(tag, base, s):
        e = OxmlElement(f"m:{tag}")
        b = OxmlElement("m:e"); b.append(mrun(base)); e.append(b)
        x = OxmlElement("m:sub" if tag == "sSub" else "m:sup"); x.append(mrun(s)); e.append(x)
        return e
    p = para("Equation")
    om = OxmlElement("m:oMath")
    for part in parts:
        if isinstance(part, str):
            om.append(mrun(part))
        else:
            om.append(script("sSub" if part[0] == "sub" else "sSup", part[1], part[2]))
    p.append(om)
    for _ in range(tabs):
        r = OxmlElement("w:r"); r.append(OxmlElement("w:tab")); p.append(r)
    p.append(mk_run(f"({label})"))
    return p


def reference_para(doc, text, url, num_id=None):
    p = para("References", ind=dict(left=576, hanging=576))
    if num_id is not None:   # the template's num 36 restarts the References numbering at 1 (used for PUBLICATION)
        npr = OxmlElement("w:numPr"); ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
        npr.append(ilvl); npr.append(nid); p.find(qn("w:pPr")).insert(1, npr)
    p.append(mk_run(text + " "))
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    h.set(qn("w:history"), "1")
    h.append(mk_run(url, underline=True))
    p.append(h)
    return p


# ------------------------------------------------------------------ template surgery
def el_text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def set_text(el, text):
    ts = list(el.iter(qn("w:t")))
    ts[0].text = text; ts[0].set(XML_SPACE, "preserve")
    for t in ts[1:]:
        t.text = ""


def style_of(el):
    ps = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


def find(body, startswith, style=None):
    for el in body:
        if el_text(el).strip().startswith(startswith) and (style is None or style_of(el) == style):
            return el
    raise KeyError(startswith)


def has_page_break(el):
    return any(b.get(qn("w:type")) == "page" for b in el.iter(qn("w:br")))


def strip_ids(el):
    for e in el.iter():
        for attr in ("{%s}paraId" % W14, "{%s}textId" % W14):
            if attr in e.attrib:
                del e.attrib[attr]
    for b in list(el.iter(qn("w:bookmarkStart"))) + list(el.iter(qn("w:bookmarkEnd"))):
        b.getparent().remove(b)
    return el


def delete_section(body, heading):
    """Remove an optional front-matter page: its Major Heading through its page break."""
    els = list(body)
    i = els.index(find(body, heading, "MajorHeading"))
    j = next(k for k in range(i, len(els)) if has_page_break(els[k]))
    for e in els[i:j + 1]:
        body.remove(e)


# ------------------------------------------------------------------ build
def build():
    prepass()
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    F = C.FRONT

    # --- title page and committee page (content controls keep their formatting)
    els = list(body)
    assert el_text(els[0]) == "TITLE OF THESIS"
    set_text(els[0], F["title"])
    set_text(els[2], F["author"])
    set_text(els[4].find(".//" + qn("w:sdt")), F["type"])
    set_text(els[8], F["degree"])
    set_text(els[11], F["department"])
    set_text(els[12].find(".//" + qn("w:sdt")), F["campus"])
    set_text(els[13].find(".//" + qn("w:sdt")), F["term"])
    slots = [(18, 19), (20, 21), (22, 23)]
    assert len(F["committee"]) == 3
    for (a, b), (name, dept) in zip(slots, F["committee"]):
        set_text(els[a], name); set_text(els[b], dept)
    for e in els[24:28]:
        assert "Add or Delete" in el_text(e) or el_text(e) == "Choose Department", el_text(e)
        body.remove(e)
    set_text(find(body, "Dr. Type the program"), F["approved_by"])
    set_text(find(body, "Dedicated to my family"), F["dedication"])

    # --- acknowledgments: replace the placeholder, keep its closing page break last
    ack = find(body, "This page is OPTIONAL")
    anchor = ack
    for text in F["acknowledgments"]:
        new = para("IndentedParagraph", text); anchor.addnext(new); anchor = new
    if has_page_break(ack):
        r = OxmlElement("w:r"); br = OxmlElement("w:br"); br.set(qn("w:type"), "page"); r.append(br); anchor.append(r)
    body.remove(ack)

    # --- optional pages not used in this thesis
    for h in ("LIST OF SCHEMES", "LIST OF SYMBOLS", "GLOSSARY"):
        delete_section(body, h)

    # --- list of abbreviations (borderless two-column table, alphabetical)
    els = list(body)
    tbl = els[els.index(find(body, "LIST OF ABBREVIATIONS", "MajorHeading")) + 1]
    assert tbl.tag == qn("w:tbl")
    rows = tbl.findall(qn("w:tr"))
    proto = deepcopy(rows[0])
    for r in rows:
        tbl.remove(r)
    for abbr, meaning in sorted(C.ABBREVIATIONS, key=lambda x: plain(x[0]).lower()):
        row = strip_ids(deepcopy(proto))
        cells = row.findall(qn("w:tc"))
        for cell, text, bold in ((cells[0], abbr, True), (cells[1], meaning, False)):
            p = cell.find(qn("w:p"))
            for r in p.findall(qn("w:r")):
                p.remove(r)
            for r in runs_for(text, bold=bold):
                p.append(r)
        tbl.append(row)

    # --- abstract
    ab = find(body, "An abstract is a concise summary")
    anchor = ab
    for text in C.ABSTRACT:
        new = para("IndentedParagraph", text); anchor.addnext(new); anchor = new
    body.remove(ab)

    # --- chapters: drop the template's three placeholder chapters and its appendix
    els = list(body)
    i0 = next(i for i, e in enumerate(els) if style_of(e) == "Heading1")
    i1 = els.index(find(body, "REFERENCES", "MajorHeading"))
    page_break = strip_ids(deepcopy(els[i0 - 1]))
    assert has_page_break(page_break) and not el_text(page_break).strip()
    for e in els[i0:i1]:
        body.remove(e)
    ref_heading = find(body, "REFERENCES", "MajorHeading")
    for chno, (title, blocks) in enumerate(C.CHAPTERS, start=1):
        out = [para("Heading1", title)]
        after_table = False
        for b in blocks:
            kind = b[0]
            if kind == "p":
                out.append(para("IndentedParagraph", NUM.resolve(b[1]), before=480 if after_table else None))
            elif kind in ("h2", "h3"):
                out.append(para("Heading2" if kind == "h2" else "Heading3", b[1]))
            elif kind == "fig":
                _, key, fname, cap, alt = b[:5]
                out.append(figure_para(doc, fname, alt, **(b[5] if len(b) > 5 else {})))
                out.append(caption("Figure", NUM.labels[("fig", key)], NUM.resolve(cap)))
            elif kind == "tab":
                _, key, d = b
                out.append(caption("Table", NUM.labels[("tab", key)], NUM.resolve(d["caption"])))
                out.append(table_el(d["header"], d["rows"], d.get("widths"), d.get("size", 20), d.get("align")))
                if d.get("note"):
                    out.append(para("NotesTable", "NOTE: " + NUM.resolve(d["note"])))
            elif kind == "eq":
                _, key, parts, tabs = b
                out.append(equation(parts, NUM.labels[("eq", key)], tabs))
            else:
                raise ValueError(kind)
            after_table = kind == "tab" and not b[2].get("note")
        out.append(deepcopy(page_break))
        for e in out:
            ref_heading.addprevious(e)

    # --- references, numbered in order of first citation
    els = list(body)
    i = els.index(ref_heading)
    j = next(k for k in range(i + 1, len(els)) if has_page_break(els[k]))
    for e in els[i + 1:j]:
        body.remove(e)
    anchor = ref_heading
    for key, _ in sorted(NUM.cites.items(), key=lambda kv: kv[1]):
        new = reference_para(doc, *REFS[key]); anchor.addnext(new); anchor = new

    # --- publication(s); the template's VITA page is not used
    pub = find(body, "PUBLICATIONS", "MajorHeading")
    els = list(body)
    i = els.index(pub)
    sect = next(k for k in range(i, len(els)) if els[k].tag == qn("w:sectPr"))
    for e in els[i + 1:sect]:
        body.remove(e)
    if C.PUBLICATIONS:
        set_text(pub, "PUBLICATION" if len(C.PUBLICATIONS) == 1 else "PUBLICATIONS")
        anchor = pub
        for text, url in C.PUBLICATIONS:
            new = reference_para(doc, text, url, num_id=36)   # same style as references, numbered from 1
            anchor.addnext(new); anchor = new
    else:
        prev = els[i - 1]
        body.remove(pub)
        if has_page_break(prev) and not el_text(prev).strip():
            body.remove(prev)

    doc.core_properties.author = F["author"]
    doc.core_properties.title = F["title_case"]
    # ask Word to recalculate every field (contents, lists, captions, cross-references) on open;
    # the template's cached placeholders are otherwise shown until the user presses F9
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        settings.insert(0, uf)
    doc.save(str(OUT))
    return doc


# ------------------------------------------------------------------ front lists
def _tab_run(text=None, tab=False):
    r = OxmlElement("w:r")
    if tab:
        r.append(OxmlElement("w:tab"))
    else:
        t = OxmlElement("w:t"); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); t.text = text; r.append(t)
    return r


def _hidden_run(child=None, text=None):
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rpr.append(OxmlElement("w:webHidden")); r.append(rpr)
    if child is not None:
        r.append(child)
    if text is not None:
        t = OxmlElement("w:t"); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); t.text = text; r.append(t)
    return r


def _entry_link(text, page, anchor):
    """One contents/list entry exactly as Word writes it: hyperlink to the heading's bookmark,
    the text, a tab, and a PAGEREF field whose cached result is the page number."""
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "Hyperlink"); rpr.append(rs); r.append(rpr)
    t = OxmlElement("w:t"); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); t.text = text; r.append(t); h.append(r)
    h.append(_hidden_run(OxmlElement("w:tab")))
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); h.append(_hidden_run(fc))
    it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); it.text = f" PAGEREF {anchor} \\h "; h.append(_hidden_run(it))
    h.append(_hidden_run())   # Word writes an empty run here; kept so the markup matches the template exactly
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); h.append(_hidden_run(fc))
    h.append(_hidden_run(text=str(page)))
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); h.append(_hidden_run(fc))
    return h


def _entry_para(ppr, text, page, anchor):
    p = OxmlElement("w:p"); p.append(deepcopy(ppr)); p.append(_entry_link(text, page, anchor))
    return p


def _bookmark(par, bid, name):
    """Bookmark the whole paragraph, as Word does for headings listed in a contents table."""
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), str(bid)); bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), str(bid))
    ppr = par.find(qn("w:pPr")); par.insert(1 if ppr is not None else 0, bs); par.append(be)


def _render_pages(docx_path):
    """LibreOffice render -> list of per-page line lists (used only for the cached page numbers)."""
    import subprocess, tempfile, pymupdf
    outdir = Path(tempfile.mkdtemp(prefix="thesis_prefill_"))
    subprocess.run(["/Applications/LibreOffice.app/Contents/MacOS/soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(docx_path)], capture_output=True, timeout=600)
    pdf = pymupdf.open(str(outdir / (docx_path.stem + ".pdf")))
    return [[l.strip() for l in pg.get_text().split("\n") if l.strip() and not l.strip().isdigit()] for pg in pdf]


def _norm(t):
    return re.sub(r"\s+", " ", t).strip().lower()


def fill_front_lists(docx_path, pages=None):
    """Write the contents, list of tables and list of figures into the cached field results so
    they show in any viewer; Word regenerates them on open (updateFields) with exact pages."""
    doc = Document(str(docx_path)); body = doc.element.body
    paras = list(body.iter(qn("w:p")))
    def ptext(p): return "".join(t.text or "" for t in p.iter(qn("w:t")))
    def pstyle(p):
        ps = p.find(qn("w:pPr") + "/" + qn("w:pStyle")); return ps.get(qn("w:val")) if ps is not None else ""
    # --- entries in document order: (level, text) for the contents; captions for the lists
    entries, tables, figures = [], [], []
    majors_on = False; chno = 0; h2 = 0; h3 = 0
    for p in paras:
        st, tx = pstyle(p), ptext(p).strip()
        if st == "MajorHeading":
            if tx == "ACKNOWLEDGMENTS": majors_on = True
            if majors_on and tx: entries.append((1, tx, tx, p))
        elif st == "Heading1":
            chno += 1; h2 = 0; entries.append((1, f"{chno}. {tx}", f"{chno}. {tx}", p))
        elif st == "Heading2":
            h2 += 1; h3 = 0; entries.append((2, f"{chno}.{h2} {tx}", f"{chno}.{h2} {tx}", p))
        elif st == "Heading3":
            h3 += 1; entries.append((3, f"{chno}.{h2}.{h3} {tx}", f"{chno}.{h2}.{h3} {tx}", p))
        elif st == "CaptionTable":
            tables.append((tx, p))
        elif st == "CaptionFigure":
            figures.append((tx, p))
    # --- page numbers from a render
    pages = pages or _render_pages(docx_path)
    def find_page(needle, first_line=False, start=0):
        n = _norm(needle)[:60]
        for i in range(start, len(pages)):
            lines = pages[i]
            if first_line:
                if lines and _norm(lines[0]).startswith(n[:40]): return i + 1
            else:
                joined = _norm(" ".join(lines))
                if n in joined: return i + 1
        return None
    body_start = next((i for i, lines in enumerate(pages) if lines and _norm(lines[0]).startswith("1. ")), 0)
    toc_rows = []; cursor = 0
    for level, text, search, par in entries:
        if level == 1:   # major headings and chapters start a page; the front lists quote them, so never search those
            pg = find_page(search, first_line=True) or find_page(search, start=max(cursor, body_start)) or (toc_rows[-1][2] if toc_rows else 1)
        else:            # subheadings: search from the current chapter's page onward
            pg = find_page(search, start=max(cursor, body_start)) or (toc_rows[-1][2] if toc_rows else 1)
        cursor = max(cursor, pg - 1)
        toc_rows.append((level, text, pg, par))
    def cap_pages(caps):
        out = []; cur = body_start   # captions are also listed in the front matter; search the body only
        for c, par in caps:
            body_txt = re.sub(r"^(Table|Figure) [\d.]+\s*", "", c)   # LibreOffice renders the number differently; match the text
            pg = find_page(body_txt[:50], start=cur) or (out[-1][1] if out else 1); cur = max(cur, pg - 1); out.append((c, pg, par))
        return out
    tab_rows, fig_rows = cap_pages(tables), cap_pages(figures)
    anchors = {}
    for k, row in enumerate([r[3] for r in toc_rows] + [r[2] for r in tab_rows] + [r[2] for r in fig_rows]):
        name = f"_Toc9{k:08d}"; _bookmark(row, 9000 + k, name); anchors[id(row)] = name
    # --- contents: replace the cached TOC paragraphs
    toc_paras = [p for p in paras if pstyle(p).startswith("TOC")]
    first = toc_paras[0]
    assert any("TOC \\o" in (t.text or "") for t in first.iter(qn("w:instrText")))
    pprs = {}
    for p in toc_paras:
        pprs.setdefault(pstyle(p), deepcopy(p.find(qn("w:pPr"))))
    for lvl in (2, 3):
        pprs.setdefault(f"TOC{lvl}", deepcopy(pprs["TOC1"]))
        ps = pprs[f"TOC{lvl}"].find(qn("w:pStyle")); ps.set(qn("w:val"), f"TOC{lvl}")
    # keep the field start of the first paragraph (begin, instr, separate), drop its cached content
    keep = []
    for child in list(first):
        keep.append(child)
        if child.tag == qn("w:r") and child.find(qn("w:fldChar")) is not None and child.find(qn("w:fldChar")).get(qn("w:fldCharType")) == "separate":
            break
    for child in list(first):
        first.remove(child)
    for child in keep: first.append(child)
    lvl, text, pg, par = toc_rows[0]
    first.append(_entry_link(text, pg, anchors[id(par)]))
    anchor = first
    for p in toc_paras[1:]:
        p.getparent().remove(p)
    for lvl, text, pg, par in toc_rows[1:]:
        e = _entry_para(pprs[f"TOC{min(lvl, 3)}"], text, pg, anchors[id(par)]); anchor.addnext(e); anchor = e
    # --- the template's LIST OF TABLES heading carries an extra, empty "TOC \c Table" field before the
    #     real one; Word fills both on update and shows the list twice, so drop the empty one
    lot_heading = next(p for p in body.iter(qn("w:p")) if pstyle(p) == "MajorHeading" and ptext(p).strip() == "LIST OF TABLES")
    runs = [r for r in lot_heading if r.tag == qn("w:r") and r.find(qn("w:fldChar")) is not None or r.tag == qn("w:r") and r.find(qn("w:instrText")) is not None]
    group = []
    for r in runs:
        group.append(r)
        fc = r.find(qn("w:fldChar"))
        if fc is not None and fc.get(qn("w:fldCharType")) == "end":
            break
    if len(group) == 4 and group[-1].find(qn("w:fldChar")).get(qn("w:fldCharType")) == "end":
        for r in group:
            lot_heading.remove(r)
    # --- list of tables: the TableofFigures paragraph that holds the cached result and the field end
    lot = next(p for p in body.iter(qn("w:p")) if pstyle(p) == "TableofFigures")
    lot_ppr = deepcopy(lot.find(qn("w:pPr")))
    end_run = next(r for r in lot.iter(qn("w:r")) if r.find(qn("w:fldChar")) is not None)
    for child in list(lot):
        lot.remove(child)
    lot.append(deepcopy(lot_ppr))
    c, pg, par = tab_rows[0]; lot.append(_entry_link(c, pg, anchors[id(par)]))
    anchor = lot
    for c, pg, par in tab_rows[1:]:
        e = _entry_para(lot_ppr, c, pg, anchors[id(par)]); anchor.addnext(e); anchor = e
    anchor.append(end_run)
    # --- list of figures: the fldSimple paragraph
    lof = next(p for p in body.iter(qn("w:p")) if p.find(qn("w:fldSimple")) is not None and "Figure" in p.find(qn("w:fldSimple")).get(qn("w:instr"), ""))
    instr = lof.find(qn("w:fldSimple")).get(qn("w:instr"))
    for child in list(lof):
        lof.remove(child)
    lof.append(deepcopy(lot_ppr))
    def fld_run(kind, text=None):
        r = OxmlElement("w:r")
        if kind == "instr":
            t = OxmlElement("w:instrText"); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); t.text = text; r.append(t)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
        return r
    lof.append(fld_run("begin")); lof.append(fld_run("instr", instr)); lof.append(fld_run("separate"))
    c, pg, par = fig_rows[0]; lof.append(_entry_link(c, pg, anchors[id(par)]))
    anchor = lof
    for c, pg, par in fig_rows[1:]:
        e = _entry_para(lot_ppr, c, pg, anchors[id(par)]); anchor.addnext(e); anchor = e
    anchor.append(fld_run("end"))
    doc.save(str(docx_path))
    return len(toc_rows), len(tab_rows), len(fig_rows)


if __name__ == "__main__":
    build()
    # two passes: the filled lists change the pagination, so fill, render the filled document,
    # rebuild and fill again with the page numbers of the filled render
    fill_front_lists(OUT)
    pages = _render_pages(OUT)
    NUM.cites.clear(); NUM.labels.clear(); _pic_id[0] = 1000
    build()
    n_toc, n_tab, n_fig = fill_front_lists(OUT, pages)
    print(f"  contents: {n_toc} entries; list of tables: {n_tab}; list of figures: {n_fig} (cached; Word refreshes on open)")
    unused = sorted(set(REFS) - set(NUM.cites))
    print(f"built {OUT.relative_to(ROOT)}: {len(C.CHAPTERS)} chapters, "
          f"{sum(1 for k in NUM.labels if k[0] == 'fig')} figures, "
          f"{sum(1 for k in NUM.labels if k[0] == 'tab')} tables, "
          f"{sum(1 for k in NUM.labels if k[0] == 'eq')} equations, {len(NUM.cites)} references cited")
    if unused:
        print("uncited reference keys:", ", ".join(unused))
